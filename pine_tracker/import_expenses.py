"""Import the monthly operating-expense workbooks into the tracker database.

Layout of the existing sheets, as read from
``Sample_Operating_Expenses_Tracking_for_September_2026.xlsx``:

    row N-1   column A/B empty, columns C..R hold "M/D description" headers
    row N     column A = property address, or column B = unit label,
              columns C..R hold the matching amounts, column S a row total
    row 59    grand total

Properties appear in column A; for the 6-unit building the property row is
followed by ``Unit 1`` .. ``Unit 6`` rows in column B. A cost on the property
row itself is a building-level cost and imports with ``unit_id`` NULL.

Scope: every property present in the workbook is imported, not a subset. The
parser is generic, so one property costs the same work as all of them, and
importing the whole sheet is what exercises all four shapes in the data at
once: property-level rows, unit-level rows, multi-item cells, and labor lines.

Usage::

    python -m pine_tracker.import_expenses \\
        --file "Sample_Operating_Expenses_Tracking_for_September_2026.xlsx" \\
        --db sqlite:///pine.db --org pine --dry-run

Re-running is safe: every line carries a ``source_ref`` such as
``spreadsheet:2026-09:Sheet1!E24#1`` and existing refs are skipped.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from .classify import classify, parse_address, parse_description
from .formula import ParseError, parse_amount_cell
from .money import fmt, to_cents
from .models import (
    Base,
    CostTreatment,
    Entity,
    ExpenseLine,
    Organization,
    Property,
    ScheduleE,
    Unit,
    Vendor,
)

# Columns C..R hold data. Column A is the property, B the unit, S the row
# total (a formula we recompute rather than trust).
FIRST_DATA_COL = 3
LAST_DATA_COL = 18
PROPERTY_COL = 1
UNIT_COL = 2
TOTAL_COL = 19

MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11,
    "december": 12,
}


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------


@dataclass
class RawEntry:
    """One cell of the sheet, before it becomes database rows."""

    sheet: str
    coordinate: str
    property_label: str
    unit_label: Optional[str]
    header: str
    formula: object
    cached_value: Optional[float]


@dataclass
class ImportStats:
    entries: int = 0
    lines_created: int = 0
    lines_skipped: int = 0
    properties_created: int = 0
    units_created: int = 0
    vendors_created: int = 0
    needs_review: int = 0
    amount_cents: int = 0
    sheet_total_cents: int = 0
    warnings: list[str] = field(default_factory=list)


def period_from_filename(path: Path) -> tuple[Optional[int], Optional[int]]:
    """``..._for_September_2026.xlsx`` into (2026, 9)."""
    stem = path.stem.lower()
    year_match = re.search(r"(20\d{2})", stem)
    year = int(year_match.group(1)) if year_match else None
    month = next((num for name, num in MONTHS.items() if name in stem), None)
    return year, month


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _header_for(header_row: int, col: int, formula_ws: Worksheet) -> str:
    """Header text for a data column, allowing a one-column offset."""
    if header_row < 1:
        return ""
    for candidate in (col, col - 1, col + 1):
        if FIRST_DATA_COL - 1 <= candidate <= LAST_DATA_COL + 1:
            text = _text(formula_ws.cell(row=header_row, column=candidate).value)
            if text and text.lower() != "total":
                return text
    return ""


def read_entries(path: Path) -> tuple[list[RawEntry], int, list[tuple[str, Optional[str]]]]:
    """Walk every sheet and return the populated cells, the sheet total, and
    the full roster of (property, unit) labels.

    The roster matters: a property with no spending this month still belongs in
    the registry, so the app opens with the whole portfolio rather than only the
    addresses that happened to have a receipt in September.

    The workbook is opened twice: once for formulas (which carry the pre-tax
    price and the tax rate) and once for cached values (the fallback, and the
    cross-check).
    """
    formula_wb = load_workbook(path, data_only=False)
    value_wb = load_workbook(path, data_only=True)

    entries: list[RawEntry] = []
    roster: list[tuple[str, Optional[str]]] = []
    seen_roster: set[tuple[str, Optional[str]]] = set()
    sheet_total_cents = 0

    for formula_ws in formula_wb.worksheets:
        value_ws = value_wb[formula_ws.title]
        current_property: Optional[str] = None

        for row in range(1, formula_ws.max_row + 1):
            property_label = _text(formula_ws.cell(row=row, column=PROPERTY_COL).value)
            unit_label = _text(formula_ws.cell(row=row, column=UNIT_COL).value)
            if property_label:
                current_property = property_label

            if current_property:
                key = (current_property, unit_label or None)
                if key not in seen_roster:
                    seen_roster.add(key)
                    roster.append(key)

            data_cols = [
                col
                for col in range(FIRST_DATA_COL, LAST_DATA_COL + 1)
                if _is_number(value_ws.cell(row=row, column=col).value)
            ]
            if not data_cols:
                continue
            if not current_property:
                continue

            row_total = value_ws.cell(row=row, column=TOTAL_COL).value
            if _is_number(row_total):
                sheet_total_cents += to_cents(row_total)

            for col in data_cols:
                cell = formula_ws.cell(row=row, column=col)
                entries.append(
                    RawEntry(
                        sheet=formula_ws.title,
                        coordinate=cell.coordinate,
                        property_label=current_property,
                        unit_label=unit_label or None,
                        header=_header_for(row - 1, col, formula_ws),
                        formula=cell.value,
                        cached_value=value_ws.cell(row=row, column=col).value,
                    )
                )

    return entries, sheet_total_cents, roster


# ---------------------------------------------------------------------------
# Writing to the database
# ---------------------------------------------------------------------------


def get_or_create_organization(session: Session, slug: str, name: str) -> Organization:
    org = session.scalar(select(Organization).where(Organization.slug == slug))
    if org is None:
        org = Organization(slug=slug, name=name)
        session.add(org)
        session.flush()
    return org


def get_or_create_entity(session: Session, org: Organization, name: str) -> Entity:
    entity = session.scalar(
        select(Entity).where(Entity.organization_id == org.id, Entity.name == name)
    )
    if entity is None:
        kind = "trust" if "trust" in name.lower() else "llc"
        entity = Entity(organization_id=org.id, name=name, kind=kind)
        session.add(entity)
        session.flush()
    return entity


def get_or_create_property(
    session: Session, org: Organization, raw_label: str, entity: Optional[Entity]
) -> tuple[Property, bool]:
    parsed = parse_address(raw_label)
    prop = session.scalar(
        select(Property).where(
            Property.organization_id == org.id, Property.street == parsed.street
        )
    )
    if prop is not None:
        return prop, False
    prop = Property(
        organization_id=org.id,
        entity_id=entity.id if entity else None,
        street=parsed.street,
        city=parsed.city,
        state=parsed.state,
        postal_code=parsed.postal_code,
        short_name=parsed.short_name,
    )
    session.add(prop)
    session.flush()
    return prop, True


def get_or_create_unit(
    session: Session, org: Organization, prop: Property, label: Optional[str]
) -> tuple[Optional[Unit], bool]:
    """Unit rows get a real unit. Property-level costs return None."""
    if not label:
        return None, False
    unit = session.scalar(
        select(Unit).where(Unit.property_id == prop.id, Unit.label == label)
    )
    if unit is not None:
        return unit, False
    unit = Unit(
        organization_id=org.id, property_id=prop.id, label=label, is_whole_property=False
    )
    session.add(unit)
    session.flush()
    return unit, True


def get_or_create_vendor(
    session: Session, org: Organization, name: Optional[str], kind: Optional[str]
) -> tuple[Optional[Vendor], bool]:
    if not name:
        return None, False
    vendor = session.scalar(
        select(Vendor).where(Vendor.organization_id == org.id, Vendor.name == name)
    )
    if vendor is not None:
        return vendor, False
    vendor = Vendor(
        organization_id=org.id,
        name=name,
        kind=kind or "retailer",
        # Labor gets 1099 tracking; retailers do not.
        track_1099=(kind == "labor"),
    )
    session.add(vendor)
    session.flush()
    return vendor, True


def import_workbook(
    session: Session,
    path: Path,
    *,
    org_slug: str = "pine",
    org_name: str = "Pine Real Estate Management",
    entity_name: Optional[str] = None,
    year: Optional[int] = None,
    month: Optional[int] = None,
) -> ImportStats:
    """Read one workbook into the database. Idempotent on ``source_ref``."""
    file_year, file_month = period_from_filename(path)
    year = year or file_year
    month = month or file_month
    if year is None:
        raise SystemExit("Could not determine the year; pass --year.")

    stats = ImportStats()
    entries, stats.sheet_total_cents, roster = read_entries(path)
    stats.entries = len(entries)

    org = get_or_create_organization(session, org_slug, org_name)
    entity = get_or_create_entity(session, org, entity_name) if entity_name else None

    # Register every address and unit the workbook mentions, whether or not it
    # had spending in this period.
    for property_label, unit_label in roster:
        prop, created = get_or_create_property(session, org, property_label, entity)
        stats.properties_created += int(created)
        _, unit_created = get_or_create_unit(session, org, prop, unit_label)
        stats.units_created += int(unit_created)

    existing_refs = set(
        session.scalars(
            select(ExpenseLine.source_ref).where(ExpenseLine.organization_id == org.id)
        )
    )
    period_tag = f"{year}-{month:02d}" if month else str(year)

    for entry in entries:
        prop, created = get_or_create_property(session, org, entry.property_label, entity)
        stats.properties_created += int(created)
        unit, created = get_or_create_unit(session, org, prop, entry.unit_label)
        stats.units_created += int(created)

        described = parse_description(entry.header or "")
        incurred = described.date_for_year(year)
        date_warning = None
        if incurred is None:
            incurred = dt.date(year, month or 1, 1)
            date_warning = "No date found in the header; defaulted to period start"
        elif month and incurred.month != month:
            date_warning = f"Header month {incurred.month} differs from workbook month {month}"

        vendor, created = get_or_create_vendor(
            session, org, described.vendor_name, described.vendor_kind
        )
        stats.vendors_created += int(created)

        # Recover the individual purchases inside the cell.
        try:
            components = parse_amount_cell(entry.formula)
            parse_warning = None
        except ParseError as exc:
            components = parse_amount_cell(entry.cached_value)
            parse_warning = f"{entry.coordinate}: {exc}; used the cached value"
            stats.warnings.append(parse_warning)

        if not components:
            continue

        base_desc = described.text or f"Imported from {entry.coordinate}"
        for index, component in enumerate(components, start=1):
            ref = f"spreadsheet:{period_tag}:{entry.sheet}!{entry.coordinate}#{index}"
            if ref in existing_refs:
                stats.lines_skipped += 1
                continue

            verdict = classify(base_desc, component.amount_cents)
            notes = [note for note in (date_warning, parse_warning, verdict.review_note) if note]
            if len(components) > 1:
                notes.append(
                    f"Cell {entry.coordinate} covered {len(components)} items "
                    f"({component.raw}); split on import"
                )

            description = base_desc
            if len(components) > 1:
                description = f"{base_desc} [item {index} of {len(components)}]"

            line = ExpenseLine(
                organization_id=org.id,
                property_id=prop.id,
                unit_id=unit.id if unit else None,
                vendor_id=vendor.id if vendor else None,
                incurred_on=incurred,
                description=description[:400],
                pretax_cents=component.pretax_cents,
                tax_cents=component.tax_cents,
                amount_cents=component.amount_cents,
                tax_rate_bp=component.tax_rate_bp,
                category=verdict.category,
                treatment=verdict.treatment,
                source="spreadsheet",
                source_ref=ref,
                needs_review=verdict.needs_review or bool(date_warning) or bool(parse_warning),
                review_note="; ".join(notes)[:300] or None,
            )
            session.add(line)
            existing_refs.add(ref)
            stats.lines_created += 1
            stats.needs_review += int(line.needs_review)
            stats.amount_cents += component.amount_cents

    session.flush()
    return stats


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def print_report(session: Session, stats: ImportStats, org_slug: str) -> None:
    org = session.scalar(select(Organization).where(Organization.slug == org_slug))
    lines = list(
        session.scalars(select(ExpenseLine).where(ExpenseLine.organization_id == org.id))
    )

    by_property: dict[str, int] = {}
    by_category: dict[str, int] = {}
    for line in lines:
        prop = session.get(Property, line.property_id)
        unit = session.get(Unit, line.unit_id) if line.unit_id else None
        key = prop.street + (f" / {unit.label}" if unit else "")
        by_property[key] = by_property.get(key, 0) + line.amount_cents
        by_category[line.category] = by_category.get(line.category, 0) + line.amount_cents

    width = max((len(k) for k in by_property), default=10)
    print("\nBy property")
    print("-" * (width + 14))
    for key in sorted(by_property):
        print(f"{key:<{width}}  {fmt(by_property[key]):>12}")
    print(f"{'TOTAL':<{width}}  {fmt(sum(by_property.values())):>12}")

    width = max((len(k) for k in by_category), default=10)
    print("\nBy Schedule E category")
    print("-" * (width + 14))
    for key in sorted(by_category):
        print(f"{key:<{width}}  {fmt(by_category[key]):>12}")

    print("\nImport summary")
    print("-" * 46)
    print(f"cells read              {stats.entries}")
    print(f"expense lines created   {stats.lines_created}")
    print(f"already imported        {stats.lines_skipped}")
    print(f"properties created      {stats.properties_created}")
    print(f"units created           {stats.units_created}")
    print(f"vendors created         {stats.vendors_created}")
    print(f"flagged for review      {stats.needs_review}")
    print(f"imported total          {fmt(stats.amount_cents)}")
    print(f"spreadsheet row totals  {fmt(stats.sheet_total_cents)}")
    delta = stats.amount_cents - stats.sheet_total_cents
    verdict = "ties out" if abs(delta) <= len(lines) else "CHECK THIS"
    print(f"difference              {fmt(delta)}  ({verdict}; rounding to whole cents)")

    if stats.warnings:
        print("\nWarnings")
        for warning in stats.warnings:
            print(f"  - {warning}")

    flagged = [line for line in lines if line.needs_review]
    if flagged:
        print("\nNeeds review")
        for line in sorted(flagged, key=lambda item: -item.amount_cents):
            print(f"  {fmt(line.amount_cents):>10}  {line.description[:60]}")
            print(f"              {line.review_note}")


def main(argv: Optional[Iterable[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--file", required=True, type=Path, nargs="+",
                        help="one or more monthly workbooks")
    parser.add_argument("--db", default="sqlite:///pine.db")
    parser.add_argument("--org", default="pine")
    parser.add_argument("--org-name", default="Pine Real Estate Management")
    parser.add_argument("--entity", default=None,
                        help="owning entity to assign, e.g. 'Kang Family Revocable Trust'")
    parser.add_argument("--year", type=int, default=None)
    parser.add_argument("--month", type=int, default=None)
    parser.add_argument("--dry-run", action="store_true",
                        help="parse and report, then roll back")
    args = parser.parse_args(list(argv) if argv is not None else None)

    engine = create_engine(args.db)
    Base.metadata.create_all(engine)

    with Session(engine) as session:
        totals = ImportStats()
        for path in args.file:
            if not path.exists():
                print(f"missing file: {path}", file=sys.stderr)
                return 2
            stats = import_workbook(
                session, path,
                org_slug=args.org, org_name=args.org_name,
                entity_name=args.entity, year=args.year, month=args.month,
            )
            for field_name in ("entries", "lines_created", "lines_skipped",
                               "properties_created", "units_created",
                               "vendors_created", "needs_review",
                               "amount_cents", "sheet_total_cents"):
                setattr(totals, field_name,
                        getattr(totals, field_name) + getattr(stats, field_name))
            totals.warnings.extend(stats.warnings)

        print_report(session, totals, args.org)
        if args.dry_run:
            session.rollback()
            print("\nDry run: nothing was saved.")
        else:
            session.commit()
            print(f"\nCommitted to {args.db}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
